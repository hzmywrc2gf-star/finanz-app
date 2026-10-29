import sys, datetime
import openpyxl
from openpyxl.styles import numbers
out = sys.argv[1]
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Haushalt'
ws.append(['Mein Haushaltsbuch 2026'])
ws.append([])
ws.append(['Datum', 'Betrag', 'Kategorie', 'Notiz', 'Art', 'Konto'])
rows = [
  (datetime.date(2026,10,1), -45.5, 'Lebensmittel', 'Wocheneinkauf Rewe', None, None),
  (datetime.date(2026,10,3), 2500, 'Gehalt', 'Oktober', None, None),
  (datetime.date(2026,10,5), -12.3, 'Yoga', 'Kurs „Rücken“ & Co', None, 'Bargeld'),
  (datetime.date(2026,10,7), 80, 'Nebeneinkünfte', 'Verkauf <Fahrrad>', 'Einnahme', None),
  ('12.10.2026', '-1.234,56 €', 'Sonstiges', 'Text-Datum und Text-Betrag', None, None),
  (datetime.date(2026,10,9), 0, 'Lebensmittel', 'Null-Betrag', None, None),
  ('kein Datum', -5, 'Lebensmittel', 'kaputtes Datum', None, None),
]
for r in rows:
    ws.append(list(r))
for row in ws.iter_rows(min_row=4, max_col=1):
    for c in row:
        if isinstance(c.value, datetime.date): c.number_format = 'DD.MM.YYYY'
ws2 = wb.create_sheet('Anderes')
ws2.append(['x', 'y'])
wb.save(out)
